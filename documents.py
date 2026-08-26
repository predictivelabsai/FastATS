"""Defensive resume text extraction without local ML models."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def extract_text(file_name: str, data: bytes) -> str:
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED:
        raise ValueError("Resume must be PDF, DOCX, TXT, or Markdown")
    try:
        if extension == ".pdf":
            import pdfplumber
            with pdfplumber.open(BytesIO(data)) as pdf:
                return "\n".join((page.extract_text() or "") for page in pdf.pages).strip()
        if extension == ".docx":
            import docx
            document = docx.Document(BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
            return "\n".join(part for part in parts if part.strip()).strip()
        return data.decode("utf-8", errors="replace").strip()
    except Exception as exc:
        raise ValueError(f"Resume text could not be extracted: {exc}") from exc
